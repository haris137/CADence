"""Excel / CSV manufacturing-record parser.

No column names or header positions are hard-coded. The parser:
  * reads every sheet (xlsx/xlsm via openpyxl with merged cells filled, xls via xlrd, csv via the csv module),
  * scores rows to find header rows (there can be several tables per sheet),
  * searches every cell for the part number (int, float, text, "Part No: 2280"...),
  * prefers matches in a part-number column, reports near matches such as "2280-A" without using them,
  * emits one Fact per non-empty cell of every matching row (plus continuation rows such as
    operation lines under a part), or every label/value pair for form-style sheets.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import logging
import math
import os
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple, Union

from .facts import Fact
from .units import normalize_unit, parse_dims, parse_number

log = logging.getLogger(__name__)

# canonical meaning -> header synonyms (lower case, punctuation stripped like _norm_header does)
CANONICAL: Dict[str, List[str]] = {
    "part_number": ["part no", "part number", "part num", "part#", "part #", "partno", "p/n", "pn", "part",
                    "item no", "item number", "item code", "item", "drawing no", "drawing number", "dwg no",
                    "dwg", "component no", "component", "part id", "part code", "part ref", "article no"],
    "description": ["description", "desc", "part name", "part description", "name", "title", "component name",
                    "designation"],
    "material": ["material", "matl", "mat'l", "mat", "material grade", "grade", "alloy", "raw material",
                 "material spec", "material specification"],
    "stock_size": ["stock size", "stock", "blank size", "raw size", "blank", "billet", "bar size",
                   "raw material size", "stock dims", "stock dimensions", "stock dim", "blank dims", "raw stock"],
    "finished_size": ["finished size", "finish size", "fin size", "fin dims", "finished dims", "finished dimensions",
                      "overall size", "overall dimensions", "overall dims", "part size", "size", "dimensions",
                      "dims", "envelope", "final size", "final dimensions", "net size"],
    "length": ["length", "len", "overall length", "oal", "l"],
    "width": ["width", "w", "breadth"],
    "thickness": ["thickness", "thk", "height", "h", "t", "depth"],
    "diameter": ["diameter", "dia", "od", "outer diameter", "outside diameter"],
    "operation": ["operation", "operation description", "op description", "process", "process description",
                  "routing step", "step description", "op desc", "operation name", "activity"],
    "op_seq": ["op seq", "op no", "op #", "op", "seq", "sequence", "op sequence", "step no", "step", "op number",
               "operation no", "operation number", "opn"],
    "machine": ["machine", "m/c", "mc", "work center", "workcenter", "work centre", "workcentre", "equipment",
                "machine tool", "cell", "machine name"],
    "setup": ["setup", "set up", "set-up", "setup no", "fixture", "clamping", "workholding"],
    "cycle_time": ["cycle time", "ct", "cycle", "run time", "machining time", "machine time", "time per part",
                   "cycle time min", "std time", "standard time"],
    "setup_time": ["setup time", "set up time", "changeover time"],
    "quantity": ["quantity", "qty", "batch qty", "batch size", "lot size", "order qty", "qty per batch"],
    "tolerance": ["tolerance", "tol", "general tolerance", "gen tol"],
    "finish": ["finish", "surface finish", "ra", "coating", "surface treatment", "plating"],
    "heat_treatment": ["heat treatment", "heat treat", "ht", "hardness", "hrc"],
    "revision": ["revision", "rev", "rev no", "version", "ver", "issue"],
    "tooling": ["tool", "tooling", "cutter", "tools", "tool list", "tool no"],
    "program": ["program", "nc program", "cnc program", "prog", "program no"],
    "units": ["units", "unit", "uom"],
    "holes": ["holes", "hole", "hole pattern", "hole size", "holes spec"],
    "notes": ["notes", "note", "remarks", "remark", "comment", "comments", "instructions", "special instructions"],
    "customer": ["customer", "client"],
    "price": ["price", "cost", "unit price", "unit cost", "rate"],
    "status": ["status", "state"],
    "date": ["date", "created", "modified", "due date"],
    "weight": ["weight", "mass", "wt"],
}

_SYN_INDEX: Dict[str, str] = {}
for _canon, _syns in CANONICAL.items():
    for _s in _syns:
        _SYN_INDEX.setdefault(_s, _canon)

PN_LABEL_RE = r"(?:part\s*(?:no\.?|number|num|#|id|code)?|p\s*/\s*n|pn|item(?:\s*(?:no\.?|code))?|dwg\.?\s*no\.?|drawing\s*(?:no\.?|number)|no\.?)"

MAX_FACTS_PER_SHEET = 220


def cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        if v.is_integer():
            return str(int(v))
        return ("%.10g" % v)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
    if isinstance(v, dt.time):
        return v.isoformat()
    return str(v).strip()


def _empty(v: Any) -> bool:
    return cell_str(v) == ""


def col_letter(c: int) -> str:
    s = ""
    c += 1
    while c:
        c, rem = divmod(c - 1, 26)
        s = chr(65 + rem) + s
    return s


def _norm_header(text: Any) -> Tuple[str, Optional[str]]:
    """'Stock Size (mm)' -> ('stock size', 'mm'); 'CT (min)' -> ('ct', 'min')."""
    s = cell_str(text)
    unit = None
    m = re.search(r"[\(\[]\s*([^)\]]{1,15})\s*[\)\]]", s)
    if m:
        unit = normalize_unit(m.group(1)) or m.group(1).strip().lower()
    core = re.sub(r"[\(\[][^)\]]*[\)\]]", " ", s).lower()
    core = core.replace("_", " ").replace(".", " ").replace(":", " ")
    core = re.sub(r"[^a-z0-9/#&' ]+", " ", core)
    core = re.sub(r"\s+", " ", core).strip()
    return core, unit


def canonical_for(header: Any) -> Optional[str]:
    core = _norm_header(header)[0].strip()
    if not core:
        return None
    if core in _SYN_INDEX:
        return _SYN_INDEX[core]
    best, best_len = None, 0
    for syn, canon in _SYN_INDEX.items():
        if len(syn) < 3:
            continue
        if re.search(r"(?<![a-z0-9])" + re.escape(syn) + r"(?![a-z0-9])", core) and len(syn) > best_len:
            best, best_len = canon, len(syn)
    return best


@dataclass
class Sheet:
    name: str
    grid: List[List[Any]]
    headers: Dict[int, Dict[int, Tuple[str, Optional[str], Optional[str]]]] = field(default_factory=dict)
    weak_header: Optional[int] = None


@dataclass
class Workbook:
    filename: str
    sheets: List[Sheet]
    warnings: List[str] = field(default_factory=list)
    kind: str = ""


@dataclass
class ExcelResult:
    facts: List[Fact] = field(default_factory=list)
    near_matches: List[Dict[str, str]] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    matched_by: Optional[str] = None
    match_locations: List[str] = field(default_factory=list)
    unknown_headers: List[str] = field(default_factory=list)


# ----------------------------------------------------------------------------------------- loading

def sniff_spreadsheet(data: bytes) -> str:
    if data[:4] == b"PK\x03\x04":
        return "xlsx"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    return "csv"


def _read_bytes(data_or_path: Union[bytes, str, os.PathLike], filename: Optional[str]) -> Tuple[bytes, str]:
    if isinstance(data_or_path, (bytes, bytearray)):
        return bytes(data_or_path), filename or "upload"
    path = os.fspath(data_or_path)
    with open(path, "rb") as fh:
        return fh.read(), filename or os.path.basename(path)


def _clean_grid(rows: List[List[Any]]) -> List[List[Any]]:
    rows = [list(r) for r in rows]
    # trim trailing empty rows
    while rows and all(_empty(v) for v in rows[-1]):
        rows.pop()
    width = 0
    for r in rows:
        for i in range(len(r) - 1, -1, -1):
            if not _empty(r[i]):
                width = max(width, i + 1)
                break
    return [(r + [None] * width)[:width] for r in rows]


def load_workbook(data_or_path: Union[bytes, str, os.PathLike], filename: Optional[str] = None) -> Workbook:
    data, filename = _read_bytes(data_or_path, filename)
    kind = sniff_spreadsheet(data)
    sheets: List[Sheet] = []
    warnings: List[str] = []
    if kind == "xlsx":
        import openpyxl  # imported here so a broken install only affects Excel
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        for ws in wb.worksheets:
            try:
                grid = [[c.value for c in row] for row in ws.iter_rows()]
                for rng in list(ws.merged_cells.ranges):
                    top = ws.cell(rng.min_row, rng.min_col).value
                    for r in range(rng.min_row, rng.max_row + 1):
                        for c in range(rng.min_col, rng.max_col + 1):
                            if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                                grid[r - 1][c - 1] = top
                sheets.append(Sheet(ws.title, _clean_grid(grid)))
            except Exception as exc:  # one bad sheet must not kill the workbook
                warnings.append(f"could not read sheet '{ws.title}': {exc}")
    elif kind == "xls":
        import pandas as pd
        try:
            frames = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=object, engine="xlrd")
        except ImportError as exc:
            raise RuntimeError("reading .xls needs the 'xlrd' package (pip install xlrd)") from exc
        for name, df in frames.items():
            grid = [[None if (isinstance(v, float) and math.isnan(v)) else v for v in row] for row in df.values.tolist()]
            sheets.append(Sheet(str(name), _clean_grid(grid)))
        warnings.append("legacy .xls: merged cells are not expanded")
    else:
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None or "\x00" in text[:2000]:
            raise ValueError("not a readable spreadsheet (unknown binary format)")
        sample = text[:5000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        sheets.append(Sheet(os.path.splitext(filename)[0] or "csv", _clean_grid(rows)))
    for sh in sheets:
        _detect_headers(sh)
    return Workbook(filename=filename, sheets=sheets, warnings=warnings, kind=kind)


# ----------------------------------------------------------------------------------------- headers

def _row_header_score(row: List[Any]) -> Tuple[int, int, int, int, bool]:
    """Returns (synonym_cells, text_cells, numeric_cells, nonempty_cells, form_like).

    form_like = the row alternates label / value ("Material | AL 6061 | Rev | B"), which is a
    form line, not a table header.
    """
    syn = text = numeric = nonempty = 0
    kinds: List[str] = []
    for v in row:
        if _empty(v):
            continue
        nonempty += 1
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            numeric += 1
            kinds.append("num")
            continue
        s = cell_str(v)
        if parse_number(s) is not None:
            numeric += 1
            kinds.append("num")
            continue
        if len(s) > 40:
            kinds.append("long")
            continue
        text += 1
        if canonical_for(s) or s.endswith(":"):
            syn += 1 if canonical_for(s) else 0
            kinds.append("label")
        else:
            kinds.append("text")
    form_like = (len(kinds) >= 2 and len(kinds) % 2 == 0
                 and all(k == "label" for k in kinds[0::2]) and all(k != "label" for k in kinds[1::2]))
    return syn, text, numeric, nonempty, form_like


def _detect_headers(sheet: Sheet) -> None:
    for r, row in enumerate(sheet.grid):
        syn, text, numeric, nonempty, form_like = _row_header_score(row)
        has_pn = any(_pn_label(v) for v in row if not _empty(v))
        if (nonempty >= 2 and syn >= 2 and numeric == 0 and not form_like
                and (syn >= 0.6 * nonempty or syn >= 3 or has_pn)):
            info = {}
            for c, v in enumerate(row):
                if _empty(v):
                    continue
                core, unit = _norm_header(v)
                info[c] = (cell_str(v), canonical_for(core), unit)
            sheet.headers[r] = info
    if not sheet.headers:
        for r, row in enumerate(sheet.grid[:15]):
            syn, text, numeric, nonempty, form_like = _row_header_score(row)
            if text >= 2 and numeric == 0 and not form_like and r + 1 < len(sheet.grid):
                sheet.weak_header = r
                break


def _header_above(sheet: Sheet, r: int) -> Optional[int]:
    best = None
    for h in sheet.headers:
        if h < r and (best is None or h > best):
            best = h
    return best


# ----------------------------------------------------------------------------------------- matching

def normalize_pn(pn: Any) -> str:
    s = cell_str(pn).strip()
    return re.sub(r"\s+", "", s)


def classify_match(value: Any, target: str) -> Optional[str]:
    """Return 'exact', 'labelled' (e.g. 'Part No: 2280'), 'near' (e.g. '2280-A') or None."""
    if _empty(value):
        return None
    t = normalize_pn(target)
    if not t:
        return None
    s = cell_str(value).strip()
    s_compact = re.sub(r"\s+", "", s)
    if s_compact.casefold() == t.casefold():
        return "exact"
    tv, sv = parse_number(t), parse_number(s_compact)
    if tv is not None and sv is not None and tv == sv and re.fullmatch(r"[0-9.+\-eE]+", s_compact):
        return "exact"  # 2280.0 vs 2280, '02280' stored as number
    if re.fullmatch(PN_LABEL_RE + r"\s*[.:#=\-]?\s*" + re.escape(t), s, re.IGNORECASE):
        return "labelled"
    if isinstance(value, float) and not value.is_integer():
        return None
    if len(s) <= 40 and re.search(r"(?<![0-9A-Za-z])" + re.escape(t) + r"(?![0-9])", s, re.IGNORECASE):
        return "near"
    return None


def _pn_label(v: Any) -> bool:
    s = cell_str(v).strip().rstrip(":").strip()
    return bool(s) and len(s) <= 25 and canonical_for(_norm_header(s)[0]) == "part_number"


def _is_label(v: Any) -> bool:
    if _empty(v) or isinstance(v, (int, float)):
        return False
    s = cell_str(v)
    if len(s) > 40 or parse_number(s) is not None:
        return False
    return s.endswith(":") or canonical_for(_norm_header(s)[0]) is not None


def _make_fact(wb: Workbook, sheet: Sheet, r: int, c: int, field_name: str, canon: Optional[str],
               unit: Optional[str], value: Any, confidence: str, row_units: Optional[str] = None,
               mode: str = "table") -> Fact:
    sval = cell_str(value)
    f = Fact(
        source_type="excel",
        source_file=wb.filename,
        location=f'sheet "{sheet.name}", row {r + 1}',
        field=field_name,
        value=sval,
        unit=unit if unit in ("mm", "in", "cm", "m", "ft", "um", "min", "s", "sec", "h", "hr", "hrs", "kg", "g") else None,
        confidence=confidence,
        category=canon or "other",
        raw=sval,
        meta={"sheet": sheet.name, "row": r + 1, "col": col_letter(c), "group": f"{sheet.name}!{r + 1}", "mode": mode},
    )
    if canon in ("stock_size", "finished_size"):
        dims = parse_dims(sval)
        if dims:
            if dims["unit"] is None:
                dims["unit"] = f.unit if f.unit in ("mm", "in", "cm", "m") else row_units
                dims["unit_assumed"] = dims["unit"] is None
            f.meta["dims"] = dims
            f.unit = dims["unit"] or f.unit
    elif canon in ("length", "width", "thickness", "diameter"):
        n = parse_number(sval)
        if n is not None:
            f.num = n
            if f.unit is None:
                f.unit = row_units
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        f.num = float(value)
    return f


def _row_units(sheet: Sheet, r: int, header: Optional[Dict[int, Tuple[str, Optional[str], Optional[str]]]]) -> Optional[str]:
    if not header:
        return None
    for c, (_, canon, _) in header.items():
        if canon == "units" and c < len(sheet.grid[r]):
            u = normalize_unit(cell_str(sheet.grid[r][c]))
            if u:
                return u
    return None


def _emit_row(wb: Workbook, sheet: Sheet, r: int, h: Optional[int], seen: set, confidence: str,
              mode: str = "table") -> List[Fact]:
    out = []
    header = sheet.headers.get(h) if h is not None else None
    weak = None
    if header is None and sheet.weak_header is not None and sheet.weak_header < r:
        weak = {c: (cell_str(v), canonical_for(_norm_header(v)[0]), _norm_header(v)[1])
                for c, v in enumerate(sheet.grid[sheet.weak_header]) if not _empty(v)}
    hdr = header or weak
    units = _row_units(sheet, r, hdr)
    for c, v in enumerate(sheet.grid[r]):
        if _empty(v) or (sheet.name, r, c) in seen:
            continue
        seen.add((sheet.name, r, c))
        if hdr and c in hdr:
            text, canon, unit = hdr[c]
        else:
            text, canon, unit = f"Column {col_letter(c)}", None, None
        out.append(_make_fact(wb, sheet, r, c, text, canon, unit, v,
                              confidence if (header or canon) else "medium", units, mode))
    return out


def _emit_form(wb: Workbook, sheet: Sheet, seen: set, confidence: str) -> List[Fact]:
    """Sheet is about one part: emit label/value pairs and any tables on it."""
    out: List[Fact] = []
    grid = sheet.grid
    for r, row in enumerate(grid):
        if r in sheet.headers:
            continue
        h = _header_above(sheet, r)
        if h is not None:
            # inside a table on the form: emit the row with its headers (stop at blank row handled by emptiness)
            if not all(_empty(v) for v in row):
                out.extend(_emit_row(wb, sheet, r, h, seen, confidence, mode="form-table"))
            continue
        c = 0
        while c < len(row):
            v = row[c]
            s = cell_str(v)
            m = re.match(r"^([A-Za-z][\w /#&'.()\-]{1,30}?)\s*[:=]\s*(.+)$", s) if s else None
            if m and canonical_for(_norm_header(m.group(1))[0]) and (sheet.name, r, c) not in seen:
                seen.add((sheet.name, r, c))
                canon = canonical_for(_norm_header(m.group(1))[0])
                out.append(_make_fact(wb, sheet, r, c, m.group(1).strip(), canon, _norm_header(m.group(1))[1],
                                      m.group(2).strip(), confidence, mode="form"))
                c += 1
                continue
            if _is_label(v) and c + 1 < len(row) and not _empty(row[c + 1]) and (sheet.name, r, c + 1) not in seen:
                label = s.rstrip(":").strip()
                core, unit = _norm_header(label)
                seen.add((sheet.name, r, c + 1))
                out.append(_make_fact(wb, sheet, r, c + 1, label, canonical_for(core), unit, row[c + 1],
                                      confidence, mode="form"))
                c += 2
                continue
            c += 1
        if len(out) >= MAX_FACTS_PER_SHEET:
            break
    return out[:MAX_FACTS_PER_SHEET]


def find_part(wb: Workbook, part_number: str) -> ExcelResult:
    res = ExcelResult(warnings=list(wb.warnings))
    target = normalize_pn(part_number)
    if not target:
        return res
    candidates: List[Tuple[int, str, Sheet, int, int, Optional[int]]] = []  # (priority, mode, sheet, r, c, header)
    near: List[Dict[str, str]] = []
    sheet_name_hits: List[Sheet] = []
    for sheet in wb.sheets:
        if classify_match(sheet.name, target) in ("exact", "labelled"):
            sheet_name_hits.append(sheet)
        for r, row in enumerate(sheet.grid):
            for c, v in enumerate(row):
                kind = classify_match(v, target)
                if kind is None:
                    continue
                if kind == "near":
                    near.append({"value": cell_str(v), "location": f'sheet "{sheet.name}", row {r + 1}, col {col_letter(c)}'})
                    continue
                h = _header_above(sheet, r)
                left = row[c - 1] if c > 0 else None
                above = sheet.grid[r - 1][c] if r > 0 else None
                canon = sheet.headers[h].get(c, ("", None, None))[1] if h is not None else None
                if h is not None and r != h and canon == "part_number":
                    candidates.append((1, "table", sheet, r, c, h))
                elif kind == "labelled" or _pn_label(left) or (h is None and _pn_label(above)):
                    candidates.append((1, "form", sheet, r, c, h))
                elif h is not None and r != h:
                    candidates.append((3, "table", sheet, r, c, h))
                elif r in sheet.headers:
                    continue  # the part number is itself sitting in a header row -> ignore
                else:
                    candidates.append((2, "row", sheet, r, c, None))

    seen: set = set()
    facts: List[Fact] = []
    if candidates:
        best = min(c[0] for c in candidates)
        chosen = [c for c in candidates if c[0] == best]
        confidence = "high" if best == 1 else ("medium" if best == 2 else "low")
        res.matched_by = "cell"
        form_done = set()
        for prio, mode, sheet, r, c, h in chosen:
            res.match_locations.append(f'sheet "{sheet.name}", row {r + 1}, col {col_letter(c)}')
            if mode == "form":
                if sheet.name not in form_done:
                    form_done.add(sheet.name)
                    facts.extend(_emit_form(wb, sheet, seen, confidence))
                continue
            facts.extend(_emit_row(wb, sheet, r, h, seen, confidence, mode=mode))
            if mode == "table":
                # continuation rows: e.g. operations listed below the part row with a blank part-number cell
                rr = r + 1
                while rr < len(sheet.grid):
                    row = sheet.grid[rr]
                    if all(_empty(v) for v in row) or rr in sheet.headers:
                        break
                    if c < len(row) and not _empty(row[c]):
                        break
                    facts.extend(_emit_row(wb, sheet, rr, h, seen, confidence, mode="continuation"))
                    rr += 1
        if best == 3:
            res.warnings.append("part number only found outside a part-number column (e.g. as a quantity); "
                                "treat these rows with caution")
    for sheet in sheet_name_hits:
        res.matched_by = res.matched_by or "sheet name"
        res.match_locations.append(f'sheet "{sheet.name}" (sheet named after the part)')
        facts.extend(_emit_form(wb, sheet, seen, "high"))
    if not facts:
        base = os.path.splitext(wb.filename)[0]
        if re.search(r"(?<![0-9A-Za-z])" + re.escape(target) + r"(?![0-9])", base, re.IGNORECASE):
            res.matched_by = "filename"
            res.match_locations.append(f"file name '{wb.filename}'")
            res.warnings.append("part number not found inside the workbook; using the whole file because its "
                                "name contains the part number")
            for sheet in wb.sheets:
                facts.extend(_emit_form(wb, sheet, seen, "medium"))
                # also emit table rows on sheets with a real header
                for r in range(len(sheet.grid)):
                    h = _header_above(sheet, r)
                    if h is not None and not all(_empty(v) for v in sheet.grid[r]):
                        facts.extend(_emit_row(wb, sheet, r, h, seen, "medium", mode="file"))
    # near matches are reported, never silently used
    uniq = []
    seen_vals = set()
    for n in near:
        key = (n["value"].casefold(), n["location"])
        if key not in seen_vals:
            seen_vals.add(key)
            uniq.append(n)
    res.near_matches = uniq[:20]
    for n in res.near_matches[:5]:
        facts.append(Fact(source_type="excel", source_file=wb.filename, location=n["location"],
                          field="Near match (NOT used)", value=n["value"], confidence="low",
                          category="near_match", raw=n["value"]))
    res.facts = _dedupe(facts)
    res.unknown_headers = sorted({f.field for f in facts if f.category == "other" and not f.field.startswith("Column ")})
    return res


ROW_LEVEL = {"operation", "op_seq", "machine", "setup", "cycle_time", "setup_time", "tooling", "program", "other"}


def _dedupe(facts: List[Fact]) -> List[Fact]:
    """Flat exports repeat part-level columns on every operation row; keep the first copy only."""
    out, seen = [], {}
    for f in facts:
        if f.category in ROW_LEVEL or f.category == "near_match":
            out.append(f)
            continue
        key = (f.field.casefold(), f.value.casefold())
        if key in seen:
            seen[key].meta["repeated_rows"] = seen[key].meta.get("repeated_rows", 1) + 1
            continue
        seen[key] = f
        out.append(f)
    return out


def list_part_numbers(wb: Workbook, limit: int = 300) -> List[str]:
    """Part numbers present in part-number columns / labels (used to suggest candidates and for 'any part')."""
    found: List[str] = []
    seen = set()

    def add(v: Any) -> None:
        s = cell_str(v).strip()
        if (s and len(s) <= 30 and re.search(r"\d", s) and s.casefold() not in seen and not _pn_label(s)
                and canonical_for(s) is None):
            seen.add(s.casefold())
            found.append(s)

    for sheet in wb.sheets:
        if re.fullmatch(r"[A-Za-z]{0,4}[-_ ]?\d{3,}[A-Za-z0-9\-]*", sheet.name.strip()):
            add(sheet.name.strip())
        hdr_rows = sorted(sheet.headers)
        for i, h in enumerate(hdr_rows):
            pn_cols = [c for c, (_, canon, _) in sheet.headers[h].items() if canon == "part_number"]
            end = hdr_rows[i + 1] if i + 1 < len(hdr_rows) else len(sheet.grid)
            for r in range(h + 1, end):
                for c in pn_cols:
                    if c < len(sheet.grid[r]):
                        add(sheet.grid[r][c])
        for r, row in enumerate(sheet.grid):
            if r in sheet.headers:
                continue
            for c, v in enumerate(row[:-1]):
                if _pn_label(v):
                    add(row[c + 1])
            for v in row:
                m = re.match(r"^" + PN_LABEL_RE + r"\s*[.:#=\-]\s*(\S+)$", cell_str(v), re.IGNORECASE)
                if m:
                    add(m.group(1))
        if len(found) >= limit:
            break
    return found[:limit]


def apply_header_mapping(facts: List[Fact], mapping: Dict[str, str]) -> int:
    """Apply an LLM-proposed header->canonical mapping (already validated against CANONICAL keys)."""
    n = 0
    for f in facts:
        if f.source_type == "excel" and f.category == "other" and f.field in mapping:
            f.category = mapping[f.field]
            f.meta["category_by"] = "llm-header-mapping"
            n += 1
    return n


def parse_excel(data_or_path: Union[bytes, str, os.PathLike], part_number: str,
                filename: Optional[str] = None) -> ExcelResult:
    """Convenience wrapper: load + find, never raises."""
    try:
        wb = load_workbook(data_or_path, filename)
    except Exception as exc:
        name = filename or (os.path.basename(os.fspath(data_or_path)) if not isinstance(data_or_path, (bytes, bytearray)) else "upload")
        return ExcelResult(warnings=[f"could not read {name}: {exc}"])
    return find_part(wb, part_number)
